"""Resume text extraction from PDF and DOCX"""
import io
import re
from pathlib import Path
from typing import Optional
from app.core.config import get_settings

settings = get_settings()


def extract_text_from_pdf(file_path: str) -> tuple[str, list[str]]:
    """
    Extract text from PDF using PyMuPDF.
    Returns (extracted_text, warnings)
    """
    warnings = []
    try:
        import fitz  # PyMuPDF
    except ImportError:
        return "", ["PyMuPDF (fitz) is not installed. Cannot extract PDF text."]

    try:
        doc = fitz.open(file_path)
        pages_text = []

        for page_num in range(len(doc)):
            page = doc[page_num]
            text = page.get_text("text")
            if text.strip():
                pages_text.append(text)
            else:
                warnings.append(f"Page {page_num + 1} appears to have no extractable text (may be scanned/image-based).")

        doc.close()

        if not pages_text:
            warnings.append("No text could be extracted from this PDF. It may be a scanned document without OCR text.")
            return "", warnings

        full_text = "\n\n".join(pages_text)
        full_text = normalize_whitespace(full_text)

        # Truncate if too long
        max_chars = settings.max_resume_text_chars
        if len(full_text) > max_chars:
            warnings.append(f"Resume text truncated from {len(full_text)} to {max_chars} characters.")
            full_text = full_text[:max_chars]

        return full_text, warnings

    except Exception as e:
        return "", [f"Error extracting PDF text: {str(e)}"]


def extract_text_from_docx(file_path: str) -> tuple[str, list[str]]:
    """
    Extract text from DOCX using python-docx.
    Returns (extracted_text, warnings)
    """
    warnings = []
    try:
        from docx import Document
    except ImportError:
        return "", ["python-docx is not installed. Cannot extract DOCX text."]

    try:
        doc = Document(file_path)
        sections = []

        for para in doc.paragraphs:
            text = para.text.strip()
            if text:
                sections.append(text)

        # Extract from tables
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    sections.append(" | ".join(row_texts))

        if not sections:
            warnings.append("No text content found in DOCX file.")
            return "", warnings

        full_text = "\n".join(sections)
        full_text = normalize_whitespace(full_text)

        max_chars = settings.max_resume_text_chars
        if len(full_text) > max_chars:
            warnings.append(f"Resume text truncated from {len(full_text)} to {max_chars} characters.")
            full_text = full_text[:max_chars]

        return full_text, warnings

    except Exception as e:
        return "", [f"Error extracting DOCX text: {str(e)}"]


def extract_resume_text(file_path: str) -> tuple[str, list[str]]:
    """
    Dispatch to the correct extractor based on file extension.
    Returns (text, warnings)
    """
    ext = Path(file_path).suffix.lower()
    if ext == ".pdf":
        return extract_text_from_pdf(file_path)
    elif ext == ".docx":
        return extract_text_from_docx(file_path)
    else:
        return "", [f"Unsupported resume format: {ext}"]


def normalize_whitespace(text: str) -> str:
    """Clean up excessive whitespace while preserving structure"""
    # Replace multiple blank lines with double newline
    text = re.sub(r"\n{3,}", "\n\n", text)
    # Replace multiple spaces with single space (but preserve newlines)
    lines = text.split("\n")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in lines]
    return "\n".join(lines).strip()
